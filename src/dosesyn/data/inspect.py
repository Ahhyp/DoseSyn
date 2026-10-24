"""M0 audit helpers: signatures, hashes, sheets, legend, schema checks."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable, List

import pandas as pd

from ..utils.config import project_root
from ..utils.io import detect_container, file_size, sha256_file
from .ids import drug_ids, cell_ids


def raw_path(cfg: Dict[str, Any], role: str) -> Path:
    return Path(cfg["files"][f"{role}_path"])


def file_entry(cfg: Dict[str, Any], role: str) -> Dict[str, Any]:
    """Describe a raw file by content only (no parsing)."""
    path = raw_path(cfg, role)
    container = detect_container(path)
    digest = sha256_file(path)
    expected = cfg.get("expected_sha256", {}).get(role)
    return {
        "role": role,
        "filename": path.name,
        "path_relative": str(path.relative_to(project_root())),
        "size_bytes": file_size(path),
        "sha256": digest,
        "sha256_expected": expected,
        "sha256_match": (expected is None) or (digest == expected),
        **container,
    }


def workbook_sheets(path: str | Path) -> List[Dict[str, Any]]:
    """List sheet dimensions using the reader appropriate for the container."""
    path = Path(path)
    info = detect_container(path)
    if info["is_ooxml"]:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [
                {
                    "name": ws.title,
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                }
                for ws in wb.worksheets
            ]
        finally:
            wb.close()
    if info["is_legacy_xls"]:
        import xlrd

        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            return [
                {
                    "name": sh.name,
                    "max_row": sh.nrows,
                    "max_column": sh.ncols,
                }
                for sh in book.sheets()
            ]
        finally:
            book.release_resources()
    raise ValueError(f"{path} is not a readable spreadsheet ({info['note']})")


def expected_columns(cfg: Dict[str, Any], role: str) -> List[str]:
    """Flatten the configured raw column names into sheet order."""
    spec = cfg["raw_columns"][role]
    cols: List[str] = []
    for key, value in spec.items():
        if key in ("sheet",):
            continue
        if key == "viability":
            cols.extend(value)
        else:
            cols.append(value)
    return cols


def check_columns(
    frame: pd.DataFrame, cfg: Dict[str, Any], role: str
) -> Dict[str, Any]:
    expected = expected_columns(cfg, role)
    actual = [str(c) for c in frame.columns]
    missing = [c for c in expected if c not in actual]
    extra = [c for c in actual if c not in expected]
    ordered = actual == expected
    return {
        "expected": expected,
        "actual": actual,
        "missing": missing,
        "extra": extra,
        "exact_order": ordered,
        "ok": not missing and not extra,
    }


def duplicate_column_names(frame: pd.DataFrame) -> List[str]:
    cols = [str(c) for c in frame.columns]
    return sorted({c for c in cols if cols.count(c) > 1})


def parse_legend(cfg: Dict[str, Any]) -> Dict[str, Any]:
    """Parse the two side-by-side field-definition blocks of the legend sheet."""
    path = raw_path(cfg, "legend")
    sheet = cfg["raw_columns"]["legend"]["sheet"]
    table = pd.read_excel(path, sheet_name=sheet, header=None, engine="openpyxl")
    blocks: Dict[str, Dict[str, str]] = {"mono": {}, "combo": {}}
    # Column 0/1 = single-agent block, column 3/4 = combination block.
    specs = {"mono": (0, 1), "combo": (3, 4)}
    for role, (key_col, desc_col) in specs.items():
        if key_col >= table.shape[1]:
            continue
        for _, row in table.iterrows():
            key = row.iloc[key_col]
            desc = row.iloc[desc_col] if desc_col < table.shape[1] else None
            if isinstance(key, str) and key.strip() and pd.notna(desc):
                clean_key = key.strip()
                clean_desc = " ".join(str(desc).split())
                if clean_key.lower() in ("batchid",) and "batchid" in blocks[role]:
                    continue
                blocks[role].setdefault(clean_key, clean_desc)
    return {
        "source_file": path.name,
        "sheet": sheet,
        "n_rows": int(table.shape[0]),
        "n_columns": int(table.shape[1]),
        "fields": blocks,
    }


def raw_entities(mono: pd.DataFrame, combo: pd.DataFrame) -> Dict[str, Any]:
    """Entity inventory shared by the manifest, data dictionary and QC."""
    cell_names = sorted(set(mono["cell_line"]) | set(combo["cell_line"]))
    drug_names = sorted(set(mono["drug_name"]) | set(combo["drugA_name"]) | set(combo["drugB_name"]))
    cids = cell_ids(cell_names)
    dids = drug_ids(drug_names)
    pairs = set()
    for a, b in zip(combo["drugA_name"], combo["drugB_name"]):
        pairs.add(tuple(sorted((a, b))))
    return {
        "cell_lines": cell_names,
        "drugs": drug_names,
        "cell_ids": cids,
        "drug_ids": dids,
        "n_cells": len(cell_names),
        "n_drugs": len(drug_names),
        "n_unordered_pairs": len(pairs),
        "pairs": sorted(" & ".join(p) for p in pairs),
        "drugs_only_as_a": sorted(set(combo["drugA_name"]) - set(combo["drugB_name"])),
        "drugs_only_as_b": sorted(set(combo["drugB_name"]) - set(combo["drugA_name"])),
    }


def missing_value_report(frame: pd.DataFrame) -> Dict[str, int]:
    return {str(c): int(frame[c].isna().sum()) for c in frame.columns}
